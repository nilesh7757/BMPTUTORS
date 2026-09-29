import openpyxl
import sqlite3
import csv
import os
import sys
import time
import pandas as pd
from clean_rules import (
    clean_location, clean_experience_role, infer_tutor_seniority, clean_qualification,
    clean_works_as, clean_gender, clean_speaks, clean_bool,
    parse_numeric_exp, parse_fee, parse_subject_and_level
)

INPUT_FILE = "/home/nilesh7757/BMPTUTORS1/finalSKAoutput.xlsx"
DB_FILE = "/home/nilesh7757/BMPTUTORS1/cleaned_tutors_data.db"
CSV_TEACHERS = "/home/nilesh7757/BMPTUTORS1/cleaned_teachers.csv"
CSV_BRIDGE = "/home/nilesh7757/BMPTUTORS1/cleaned_teacher_subjects.csv"
PARQUET_TEACHERS = "/home/nilesh7757/BMPTUTORS1/cleaned_teachers.parquet"
PARQUET_BRIDGE = "/home/nilesh7757/BMPTUTORS1/cleaned_teacher_subjects.parquet"

TIER_1_CITIES = {"Delhi", "Mumbai", "Bengaluru", "Bangalore", "Hyderabad", "Chennai", "Kolkata", "Pune", "Ahmedabad"}

def get_city_tier(city, state):
    if not city or city == "Not Specified":
        return "Tier 3 / Rural"
    if city in TIER_1_CITIES or state in ["Delhi", "Chandigarh"]:
        return "Tier 1 (Metro)"
    elif city in ["Jaipur", "Lucknow", "Kanpur", "Nagpur", "Indore", "Thane", "Bhopal", "Visakhapatnam", "Patna", "Vadodara", "Ghaziabad", "Ludhiana", "Agra", "Nashik", "Faridabad", "Meerut", "Rajkot", "Varanasi", "Srinagar", "Aurangabad", "Dhanbad", "Amritsar", "Navi Mumbai", "Allahabad", "Ranchi", "Howrah", "Coimbatore", "Jabalpur", "Gwalior", "Vijayawada", "Jodhpur", "Madurai", "Raipur", "Kota", "Guwahati", "Chandigarh", "Solapur", "Hubli", "Mysore", "Tiruchirappalli", "Bareilly", "Aligarh", "Tiruppur", "Gurgaon", "Gurugram", "Noida", "Dehradun"]:
        return "Tier 2 (Urban)"
    return "Tier 3 / Semi-Urban"

def run_pipeline():
    start_time = time.time()
    print("Starting Comprehensive Contextual Data Cleaning & Bridge Table Pipeline...")

    # 1. Load validation data
    print("Loading validation status...")
    wb_val = openpyxl.load_workbook(INPUT_FILE, read_only=True)
    val_sheet = wb_val["validation"]
    validation_map = {}
    for i, row in enumerate(val_sheet.iter_rows(values_only=True)):
        if i == 0 or not row[0]:
            continue
        url = str(row[0]).strip()
        is_active = row[1]
        reason = row[2] if len(row) > 2 else None
        validation_map[url] = (is_active, reason)
    wb_val.close()
    print(f"Loaded validation data for {len(validation_map)} URLs.")

    # 2. Setup SQLite database
    if os.path.exists(DB_FILE):
        os.remove(DB_FILE)
    conn = sqlite3.connect(DB_FILE)
    cur = conn.cursor()

    cur.execute("""
    CREATE TABLE teachers (
        tutor_id TEXT PRIMARY KEY,
        profile_url TEXT,
        teacher_name TEXT,
        teacher_hookline TEXT,
        description TEXT,
        location_raw TEXT,
        state TEXT,
        region TEXT,
        city TEXT,
        city_tier TEXT,
        gender TEXT,
        works_as TEXT,
        experience_raw TEXT,
        experience_role TEXT,
        seniority_tier TEXT,
        total_teaching_exp_years REAL,
        online_teaching_exp_years REAL,
        education_raw TEXT,
        highest_qualification TEXT,
        has_education_degree INTEGER,
        speaks_languages TEXT,
        num_languages INTEGER,
        teaches_online INTEGER,
        teaches_at_student_home INTEGER,
        can_travel INTEGER,
        homework_help INTEGER,
        fee_raw TEXT,
        fee_unit TEXT,
        hourly_fee_min REAL,
        hourly_fee_max REAL,
        hourly_fee_avg REAL,
        registered_date TEXT,
        last_login_date TEXT,
        is_active INTEGER,
        validation_reason TEXT,
        source_tab TEXT
    )
    """)

    cur.execute("""
    CREATE TABLE teacher_subjects (
        bridge_id INTEGER PRIMARY KEY AUTOINCREMENT,
        tutor_id TEXT,
        subject_index INTEGER,
        raw_subject_string TEXT,
        canonical_subject TEXT,
        subject_category TEXT,
        level_category TEXT,
        competitive_exam TEXT,
        FOREIGN KEY (tutor_id) REFERENCES teachers(tutor_id)
    )
    """)

    cur.execute("CREATE INDEX idx_teach_state ON teachers(state)")
    cur.execute("CREATE INDEX idx_teach_city ON teachers(city)")
    cur.execute("CREATE INDEX idx_teach_region ON teachers(region)")
    cur.execute("CREATE INDEX idx_teach_qual ON teachers(highest_qualification)")
    cur.execute("CREATE INDEX idx_teach_gender ON teachers(gender)")
    cur.execute("CREATE INDEX idx_teach_tier ON teachers(city_tier)")
    cur.execute("CREATE INDEX idx_teach_seniority ON teachers(seniority_tier)")

    cur.execute("CREATE INDEX idx_bridge_tutor ON teacher_subjects(tutor_id)")
    cur.execute("CREATE INDEX idx_bridge_category ON teacher_subjects(subject_category)")
    cur.execute("CREATE INDEX idx_bridge_canonical ON teacher_subjects(canonical_subject)")
    cur.execute("CREATE INDEX idx_bridge_level ON teacher_subjects(level_category)")
    cur.execute("CREATE INDEX idx_bridge_exam ON teacher_subjects(competitive_exam)")

    teacher_cols = [
        "tutor_id", "profile_url", "teacher_name", "teacher_hookline", "description",
        "location_raw", "state", "region", "city", "city_tier", "gender", "works_as",
        "experience_raw", "experience_role", "seniority_tier", "total_teaching_exp_years", "online_teaching_exp_years",
        "education_raw", "highest_qualification", "has_education_degree",
        "speaks_languages", "num_languages", "teaches_online", "teaches_at_student_home",
        "can_travel", "homework_help", "fee_raw", "fee_unit",
        "hourly_fee_min", "hourly_fee_max", "hourly_fee_avg",
        "registered_date", "last_login_date", "is_active", "validation_reason", "source_tab"
    ]

    bridge_cols = [
        "tutor_id", "subject_index", "raw_subject_string",
        "canonical_subject", "subject_category", "level_category", "competitive_exam"
    ]

    f_teach = open(CSV_TEACHERS, "w", newline="", encoding="utf-8")
    writer_teach = csv.writer(f_teach)
    writer_teach.writerow(teacher_cols)

    f_bridge = open(CSV_BRIDGE, "w", newline="", encoding="utf-8")
    writer_bridge = csv.writer(f_bridge)
    writer_bridge.writerow(bridge_cols)

    tutor_counter = 0
    seen_urls = set()
    total_subjects_count = 0

    batch_teachers = []
    batch_bridge = []
    BATCH_SIZE = 5000

    wb = openpyxl.load_workbook(INPUT_FILE, read_only=True)

    # 3. Process 'scraped' tab
    ws_scraped = wb["scraped"]
    rows_scraped = ws_scraped.iter_rows(values_only=True)
    header_scraped = next(rows_scraped)

    level_col_map = []
    for col_idx in range(1, 28):
        s_col = 21 + (col_idx * 2)
        l_col = 22 + (col_idx * 2)
        if s_col < len(header_scraped):
            level_col_map.append((s_col, l_col if l_col < len(header_scraped) else None))

    print(f"Found {len(level_col_map)} paired subject/level column slots in 'scraped' tab.")
    print("Processing 'scraped' tab...")

    for row_idx, row in enumerate(rows_scraped):
        url = row[0]
        if not url or not str(url).strip():
            continue
        url = str(url).strip()
        if url in seen_urls:
            continue
        seen_urls.add(url)
        
        tutor_counter += 1
        tutor_id = f"T{tutor_counter}"

        subjects_raw = row[1]
        exp_raw = row[2]
        edu_raw = row[3]
        name = row[4]
        hookline = row[5]
        desc = row[6]
        loc_raw = row[7]
        can_travel_raw = row[8] if len(row) > 8 else None
        last_login_raw = row[9] if len(row) > 9 else None
        registered_raw = row[10] if len(row) > 10 else None
        total_exp_raw = row[11] if len(row) > 11 else None
        teaches_online_raw = row[12] if len(row) > 12 else None
        online_exp_raw = row[13] if len(row) > 13 else None
        teaches_home_raw = row[14] if len(row) > 14 else None
        hw_help_raw = row[15] if len(row) > 15 else None
        gender_raw = row[16] if len(row) > 16 else None
        works_as_raw = row[17] if len(row) > 17 else None
        speaks_raw = row[18] if len(row) > 18 else None
        fee_raw = row[19] if len(row) > 19 else None

        # Clean teacher metadata
        state, region, loc_clean, city = clean_location(loc_raw)
        city_tier = get_city_tier(city, state)
        exp_role = clean_experience_role(exp_raw)
        highest_qual, has_edu_deg = clean_qualification(edu_raw)
        gender_clean = clean_gender(gender_raw)
        works_as_clean = clean_works_as(works_as_raw)
        langs_clean, num_langs = clean_speaks(speaks_raw)
        
        total_exp_yrs = parse_numeric_exp(total_exp_raw)
        online_exp_yrs = parse_numeric_exp(online_exp_raw)
        seniority_tier = infer_tutor_seniority(total_exp_yrs, highest_qual, exp_role, hookline, desc)
        
        teaches_online = clean_bool(teaches_online_raw)
        teaches_home = clean_bool(teaches_home_raw)
        can_travel = clean_bool(can_travel_raw)
        hw_help = clean_bool(hw_help_raw)
        
        fee_unit, fee_min, fee_max, fee_avg, hourly_fee = parse_fee(fee_raw)
        
        last_login = str(last_login_raw).replace("Last login:", "").strip() if last_login_raw else None
        registered = str(registered_raw).replace("Registered:", "").strip() if registered_raw else None
        
        val_info = validation_map.get(url, (None, None))
        is_active = val_info[0]
        val_reason = val_info[1]

        teacher_record = (
            tutor_id, url, name, hookline, desc,
            loc_clean, state, region, city, city_tier, gender_clean, works_as_clean,
            exp_raw, exp_role, seniority_tier, total_exp_yrs, online_exp_yrs,
            edu_raw, highest_qual, has_edu_deg,
            langs_clean, num_langs, teaches_online, teaches_home,
            can_travel, hw_help, fee_raw, fee_unit,
            fee_min if fee_unit == "hour" else None,
            fee_max if fee_unit == "hour" else None,
            hourly_fee,
            registered, last_login, is_active, val_reason, "scraped"
        )

        batch_teachers.append(teacher_record)
        writer_teach.writerow(teacher_record)

        # Process paired subject & level slots
        subj_idx = 0
        found_any_paired_subject = False
        for s_idx, l_idx in level_col_map:
            s_val = row[s_idx]
            l_val = row[l_idx] if l_idx is not None else None
            if s_val and str(s_val).strip():
                parsed = parse_subject_and_level(s_val, l_val, hookline, desc, highest_qual, exp_role)
                if parsed:
                    found_any_paired_subject = True
                    subj_idx += 1
                    total_subjects_count += 1
                    bridge_record = (
                        tutor_id,
                        subj_idx,
                        parsed["raw_subject"],
                        parsed["canonical_subject"],
                        parsed["subject_category"],
                        parsed["level_category"],
                        parsed["competitive_exam"]
                    )
                    batch_bridge.append(bridge_record)
                    writer_bridge.writerow(bridge_record)

        if not found_any_paired_subject and subjects_raw and str(subjects_raw).strip():
            subj_str = str(subjects_raw)
            tokens = subj_str.split("::") if "::" in subj_str else subj_str.split(",")
            for tok in tokens:
                parsed = parse_subject_and_level(tok, None, hookline, desc, highest_qual, exp_role)
                if parsed:
                    subj_idx += 1
                    total_subjects_count += 1
                    bridge_record = (
                        tutor_id,
                        subj_idx,
                        parsed["raw_subject"],
                        parsed["canonical_subject"],
                        parsed["subject_category"],
                        parsed["level_category"],
                        parsed["competitive_exam"]
                    )
                    batch_bridge.append(bridge_record)
                    writer_bridge.writerow(bridge_record)

        if len(batch_teachers) >= BATCH_SIZE:
            cur.executemany("INSERT INTO teachers VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)", batch_teachers)
            cur.executemany("INSERT INTO teacher_subjects (tutor_id, subject_index, raw_subject_string, canonical_subject, subject_category, level_category, competitive_exam) VALUES (?,?,?,?,?,?,?)", batch_bridge)
            conn.commit()
            batch_teachers.clear()
            batch_bridge.clear()
            print(f"Processed {tutor_counter} tutors... ({total_subjects_count} bridge records)")

    # 4. Process 'not scraped' tab
    ws_not_scraped = wb["not scraped"]
    rows_not_scraped = ws_not_scraped.iter_rows(values_only=True)
    header_not_scraped = next(rows_not_scraped)

    print("Processing 'not scraped' tab...")
    for row_idx, row in enumerate(rows_not_scraped):
        url = row[0]
        if not url or not str(url).strip():
            continue
        url = str(url).strip()
        if url in seen_urls:
            continue
        seen_urls.add(url)
        
        tutor_counter += 1
        tutor_id = f"T{tutor_counter}"

        subjects_raw = row[1]
        exp_raw = row[2]
        edu_raw = row[3]
        name = row[4]
        hookline = row[5]
        desc = row[6]
        loc_raw = row[7]
        can_travel_raw = row[8] if len(row) > 8 else None
        last_login_raw = row[9] if len(row) > 9 else None
        registered_raw = row[10] if len(row) > 10 else None
        total_exp_raw = row[11] if len(row) > 11 else None
        teaches_online_raw = row[12] if len(row) > 12 else None
        online_exp_raw = row[13] if len(row) > 13 else None
        teaches_home_raw = row[14] if len(row) > 14 else None
        hw_help_raw = row[15] if len(row) > 15 else None
        gender_raw = row[16] if len(row) > 16 else None
        works_as_raw = row[17] if len(row) > 17 else None
        speaks_raw = row[18] if len(row) > 18 else None
        fee_raw = row[19] if len(row) > 19 else None

        state, region, loc_clean, city = clean_location(loc_raw)
        city_tier = get_city_tier(city, state)
        exp_role = clean_experience_role(exp_raw)
        highest_qual, has_edu_deg = clean_qualification(edu_raw)
        gender_clean = clean_gender(gender_raw)
        works_as_clean = clean_works_as(works_as_raw)
        langs_clean, num_langs = clean_speaks(speaks_raw)
        
        total_exp_yrs = parse_numeric_exp(total_exp_raw)
        online_exp_yrs = parse_numeric_exp(online_exp_raw)
        seniority_tier = infer_tutor_seniority(total_exp_yrs, highest_qual, exp_role, hookline, desc)
        
        teaches_online = clean_bool(teaches_online_raw)
        teaches_home = clean_bool(teaches_home_raw)
        can_travel = clean_bool(can_travel_raw)
        hw_help = clean_bool(hw_help_raw)
        
        fee_unit, fee_min, fee_max, fee_avg, hourly_fee = parse_fee(fee_raw)
        
        last_login = str(last_login_raw).replace("Last login:", "").strip() if last_login_raw else None
        registered = str(registered_raw).replace("Registered:", "").strip() if registered_raw else None
        
        val_info = validation_map.get(url, (None, None))
        is_active = val_info[0]
        val_reason = val_info[1]

        teacher_record = (
            tutor_id, url, name, hookline, desc,
            loc_clean, state, region, city, city_tier, gender_clean, works_as_clean,
            exp_raw, exp_role, seniority_tier, total_exp_yrs, online_exp_yrs,
            edu_raw, highest_qual, has_edu_deg,
            langs_clean, num_langs, teaches_online, teaches_home,
            can_travel, hw_help, fee_raw, fee_unit,
            fee_min if fee_unit == "hour" else None,
            fee_max if fee_unit == "hour" else None,
            hourly_fee,
            registered, last_login, is_active, val_reason, "not scraped"
        )

        batch_teachers.append(teacher_record)
        writer_teach.writerow(teacher_record)

        # Subjects processing
        if subjects_raw and str(subjects_raw).strip():
            subj_str = str(subjects_raw)
            tokens = subj_str.split("::") if "::" in subj_str else subj_str.split(",")
            subj_idx = 0
            for tok in tokens:
                parsed = parse_subject_and_level(tok, None, hookline, desc, highest_qual, exp_role)
                if parsed:
                    subj_idx += 1
                    total_subjects_count += 1
                    bridge_record = (
                        tutor_id,
                        subj_idx,
                        parsed["raw_subject"],
                        parsed["canonical_subject"],
                        parsed["subject_category"],
                        parsed["level_category"],
                        parsed["competitive_exam"]
                    )
                    batch_bridge.append(bridge_record)
                    writer_bridge.writerow(bridge_record)

        if len(batch_teachers) >= BATCH_SIZE:
            cur.executemany("INSERT INTO teachers VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)", batch_teachers)
            cur.executemany("INSERT INTO teacher_subjects (tutor_id, subject_index, raw_subject_string, canonical_subject, subject_category, level_category, competitive_exam) VALUES (?,?,?,?,?,?,?)", batch_bridge)
            conn.commit()
            batch_teachers.clear()
            batch_bridge.clear()
            print(f"Processed {tutor_counter} tutors... ({total_subjects_count} bridge records)")

    # Final flush
    if batch_teachers:
        cur.executemany("INSERT INTO teachers VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)", batch_teachers)
        conn.commit()
        batch_teachers.clear()
    if batch_bridge:
        cur.executemany("INSERT INTO teacher_subjects (tutor_id, subject_index, raw_subject_string, canonical_subject, subject_category, level_category, competitive_exam) VALUES (?,?,?,?,?,?,?)", batch_bridge)
        conn.commit()
        batch_bridge.clear()

    wb.close()
    f_teach.close()
    f_bridge.close()
    conn.close()

    print("Converting to Parquet for high-speed analysis...")
    df_t = pd.read_csv(CSV_TEACHERS, low_memory=False)
    df_t.to_parquet(PARQUET_TEACHERS, index=False)
    df_b = pd.read_csv(CSV_BRIDGE, low_memory=False)
    df_b.to_parquet(PARQUET_BRIDGE, index=False)

    elapsed = round(time.time() - start_time, 2)
    print("\n========================================================")
    print(f"PIPELINE COMPLETE in {elapsed} seconds!")
    print(f"Total Unique Tutors Cleaned: {tutor_counter}")
    print(f"Total Bridge Subject Records: {total_subjects_count}")
    print(f"Saved to SQLite: {DB_FILE}")
    print(f"Saved to CSV: {CSV_TEACHERS}")
    print(f"Saved to CSV: {CSV_BRIDGE}")
    print(f"Saved to Parquet: {PARQUET_TEACHERS}")
    print(f"Saved to Parquet: {PARQUET_BRIDGE}")
    print("========================================================")

if __name__ == "__main__":
    run_pipeline()
